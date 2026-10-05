from __future__ import annotations

import hashlib
import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}

HEADER_ALIASES = {
    "order_code": {
        "nr zl", "nr zlecenia", "zlecenie", "zl", "numer zl", "nr zlecenia produkcyjnego",
    },
    "symbol": {
        "symbol", "kod", "indeks", "nr produktu", "numer produktu", "symbol produktu",
    },
    "product": {
        "produkt", "produkt / nazwa", "produkt/nazwa", "wyrób", "wyrob",
    },
    "name": {
        "nazwa", "nazwa produktu", "opis", "nazwa wyrobu",
    },
    "quantity": {
        "ilosc", "ilość", "qty", "szt", "szt.", "ilosc szt", "ilość szt",
    },
    "shipping": {
        "wysylka", "wysyłka", "termin", "data wysylki", "data wysyłki", "termin wysylki", "termin wysyłki",
    },
    "ral": {
        "ral", "kolor", "kolor ral",
    },
}


@dataclass(frozen=True)
class SnapshotInfo:
    path: Path
    source_name: str
    sha256: str
    size_bytes: int
    created_at: str


@dataclass(frozen=True)
class ParsedPlan:
    sheet_name: str
    header_row: int
    rows: list[dict]


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().casefold()
    text = re.sub(r"\s+", " ", text)
    return text


def safe_snapshot(source: Path, snapshot_dir: Path) -> SnapshotInfo:
    """Kopiuje źródło bajt po bajcie i zamyka je przed jakimkolwiek parsowaniem."""
    source = Path(source)
    snapshot_dir = Path(snapshot_dir)

    if not source.is_file():
        raise FileNotFoundError(f"Nie znaleziono pliku planu: {source}")

    extension = source.suffix.casefold()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Obsługiwane są pliki Excel .xlsx i .xlsm.")

    snapshot_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    temp_path = snapshot_dir / f".snapshot-{stamp}.tmp"

    digest = hashlib.sha256()
    size_bytes = 0

    # Krytyczna zasada: uchwyt do źródła istnieje wyłącznie w tym bloku.
    with source.open("rb") as source_handle, temp_path.open("wb") as target_handle:
        while True:
            chunk = source_handle.read(1024 * 1024)
            if not chunk:
                break
            target_handle.write(chunk)
            digest.update(chunk)
            size_bytes += len(chunk)
        target_handle.flush()
        os.fsync(target_handle.fileno())

    sha256 = digest.hexdigest()
    final_path = snapshot_dir / f"plan-{stamp}-{sha256[:10]}{extension}"
    os.replace(temp_path, final_path)

    return SnapshotInfo(
        path=final_path,
        source_name=source.name,
        sha256=sha256,
        size_bytes=size_bytes,
        created_at=datetime.now(timezone.utc).isoformat(),
    )


def _detect_header_row(sheet, max_scan_rows: int = 80) -> tuple[int, dict[str, int]]:
    best_row = 0
    best_mapping: dict[str, int] = {}
    best_score = 0

    for row_no, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row, start=1):
            normalized = _normalize_header(value)
            if not normalized:
                continue
            for field, aliases in HEADER_ALIASES.items():
                if normalized in aliases and field not in mapping:
                    mapping[field] = column_index

        score = len(mapping)
        if "quantity" in mapping:
            score += 1
        if "symbol" in mapping or "product" in mapping:
            score += 2
        if score > best_score:
            best_score = score
            best_row = row_no
            best_mapping = mapping

    if best_score < 3 or not (
        "symbol" in best_mapping or "product" in best_mapping
    ):
        raise ValueError(
            "Nie rozpoznano nagłówków planu. Potrzebny jest co najmniej produkt/symbol "
            "oraz dodatkowe kolumny, np. ilość lub numer ZL."
        )

    return best_row, best_mapping


def _cell(row: tuple, column_index: int | None) -> object:
    if not column_index or column_index < 1 or column_index > len(row):
        return None
    return row[column_index - 1]


def _text(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return str(value).strip()


def _number(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _split_product(product_text: str) -> tuple[str, str]:
    product_text = product_text.strip()
    if not product_text:
        return "", ""
    parts = product_text.split(maxsplit=1)
    first = parts[0]
    if re.match(r"^[A-Za-z0-9]+(?:[.\-_/][A-Za-z0-9]+)+$", first):
        return first, parts[1].strip() if len(parts) > 1 else ""
    return "", product_text


def read_plan_snapshot(snapshot_path: Path) -> ParsedPlan:
    """Parsuje wyłącznie lokalny snapshot; nigdy plik źródłowy."""
    snapshot_path = Path(snapshot_path)
    workbook = load_workbook(
        filename=snapshot_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        sheet = workbook.active
        header_row, mapping = _detect_header_row(sheet)

        result: list[dict] = []
        last_order_code = ""
        occurrence: dict[str, int] = {}

        for row_no, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            order_code = _text(_cell(row, mapping.get("order_code")))
            if order_code:
                last_order_code = order_code
            else:
                order_code = last_order_code

            symbol = _text(_cell(row, mapping.get("symbol")))
            name = _text(_cell(row, mapping.get("name")))
            product_text = _text(_cell(row, mapping.get("product")))
            if product_text and not symbol:
                inferred_symbol, inferred_name = _split_product(product_text)
                symbol = inferred_symbol
                if not name:
                    name = inferred_name
            elif product_text and not name:
                name = product_text

            quantity = _number(_cell(row, mapping.get("quantity")))
            shipping = _text(_cell(row, mapping.get("shipping")))
            ral = _text(_cell(row, mapping.get("ral")))

            if not any((order_code, symbol, name, quantity is not None, shipping, ral)):
                continue
            if not symbol and not name:
                continue

            base_key = "|".join(
                part.casefold().strip()
                for part in (order_code, symbol, name)
            )
            occurrence[base_key] = occurrence.get(base_key, 0) + 1
            row_key = f"{base_key}#{occurrence[base_key]}"

            result.append(
                {
                    "row_key": row_key,
                    "row_no": row_no,
                    "order_code": order_code,
                    "symbol": symbol,
                    "name": name,
                    "quantity": quantity,
                    "shipping": shipping,
                    "ral": ral,
                }
            )

        return ParsedPlan(
            sheet_name=str(sheet.title),
            header_row=header_row,
            rows=result,
        )
    finally:
        workbook.close()


def compare_plan_rows(previous_rows: Iterable[dict], current_rows: Iterable[dict]) -> dict:
    previous = {str(row["row_key"]): dict(row) for row in previous_rows}
    current = {str(row["row_key"]): dict(row) for row in current_rows}

    added = [current[key] for key in current.keys() - previous.keys()]
    removed = [previous[key] for key in previous.keys() - current.keys()]

    fields = ("order_code", "symbol", "name", "quantity", "shipping", "ral")
    changed: list[dict] = []
    for key in current.keys() & previous.keys():
        before = previous[key]
        after = current[key]
        delta = {
            field: {"before": before.get(field), "after": after.get(field)}
            for field in fields
            if before.get(field) != after.get(field)
        }
        if delta:
            changed.append(
                {
                    "row_key": key,
                    "before": before,
                    "after": after,
                    "changes": delta,
                }
            )

    return {
        "added": sorted(added, key=lambda row: int(row.get("row_no", 0))),
        "removed": sorted(removed, key=lambda row: int(row.get("row_no", 0))),
        "changed": sorted(
            changed,
            key=lambda row: int(row["after"].get("row_no", 0)),
        ),
    }
